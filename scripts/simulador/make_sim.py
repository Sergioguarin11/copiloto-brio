"""Genera docs/index.html (GitHub Pages) leyendo anexo/03_Anexo_Calculos.xlsx, la misma fuente auditable.

Uso: python3 scripts/simulador/make_sim.py  (requiere openpyxl)
"""
import openpyxl, json, pathlib
REPO = pathlib.Path(__file__).resolve().parents[2]
wb = openpyxl.load_workbook(REPO / 'anexo' / '03_Anexo_Calculos.xlsx', data_only=True)

def tabla(hoja, fila_inicio, prefijo):
    out = []
    for r in wb[hoja].iter_rows(min_row=fila_inicio, values_only=True):
        if r[0] and str(r[0]).startswith(prefijo):
            out.append(list(r))
    return out

def limpio(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v

datos = [[limpio(c) for c in r[:7]] for r in tabla('Datos', 4, 'DAT-')]
supuestos = [[limpio(c) for c in r[:6]] for r in tabla('Supuestos', 4, 'S-')]
supuestos.sort(key=lambda r: int(r[0][2:]))
decisiones = [[limpio(c) for c in r[:5]] for r in tabla('Decisiones', 4, 'D-')]
decisiones.sort(key=lambda r: int(r[0][2:]))
pbis = []
for r in tabla('PBIs', 4, 'PBI-'):
    pbis.append({
        'id': r[0], 'frente': r[1], 'area': r[2], 'titulo': r[3], 'para': r[4], 'criterio': r[5],
        'ing': [r[7], r[8], r[9]], 'dat': [r[10], r[11], r[12]], 'seg': [r[13], r[14], r[15]],
        'exp': [r[16], r[17], r[18]], 'incluido': r[19], 'nota': r[20] or ''})
portafolio = [[limpio(c) for c in r[:9]] for r in tabla('Portafolio', 4, 'IA-')]
fondo = []
for r in wb['Fondo_Externo'].iter_rows(min_row=5, max_row=8, values_only=True):
    fondo.append({'uso': r[0], 'horas': r[1], 'tarifa': r[2], 'costo': r[3], 'incluido': r[4], 'porque': r[6]})

dv = {r[0]: r[2] for r in datos}
sv = {r[0]: r[2] for r in supuestos}
cap = wb['Capacidad']
costo_ia01 = cap['G20'].value
assert abs(costo_ia01 - 17925000) < 1, costo_ia01

D = {
    'datos': datos, 'supuestos': supuestos, 'decisiones': decisiones, 'pbis': pbis,
    'portafolio': portafolio, 'fondo': fondo,
    'param': {
        'analistasObs': dv['DAT-02'], 'analistasTot': dv['DAT-01'], 'hSemana': dv['DAT-03'],
        'semanas': dv['DAT-04'], 'costoHora': dv['DAT-05'], 'costoIA01': costo_ia01,
        'meta': sv['S-16'], 'umbral': sv['S-17'], 'promesa': dv['DAT-33'], 'captura': sv['S-12'],
        'disp': {'ing': dv['DAT-23'], 'dat': dv['DAT-24'], 'seg': dv['DAT-25'], 'exp': dv['DAT-26']},
        'reservaMin': sv['S-14'], 'techoFondo': dv['DAT-30'],
        'tarifas': {'ing': dv['DAT-27'], 'dat': dv['DAT-28'], 'seg': dv['DAT-29']},
    },
}
html = (REPO / 'scripts/simulador/plantilla.html').read_text()
cuerpo = html.replace('/*__DATA__*/null', json.dumps(D, ensure_ascii=False))
# Página completa para GitHub Pages (la plantilla es solo el cuerpo)
out = ('<!doctype html>\n<html lang="es">\n<head>\n<meta charset="utf-8">\n'
       '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
       '<meta name="description" content="Simulador del plan de IA a 90 días de Almena Foods (prueba técnica, datos sintéticos)">\n'
       + cuerpo + '\n</html>\n')
for nombre in ('index.html', 'simulador.html'):
    (REPO / 'docs' / nombre).write_text(out)
print('ok', len(out), 'bytes;', len(datos), 'datos', len(supuestos), 'supuestos', len(decisiones), 'decisiones', len(pbis), 'pbis', len(portafolio), 'pedidos')
